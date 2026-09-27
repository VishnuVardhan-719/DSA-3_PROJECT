"""Build the editable academic project report from verified repository facts."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Contract_Compliance_Final_Project_Report.docx"
NAVY = "173A32"
GREEN = "2D7358"
PALE = "EDF4F0"
GRAY = "E5E9E7"
TEXT = RGBColor(31, 41, 37)


def set_cell_fill(cell, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = margins.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, end])


def add_table(doc, headers, rows, widths=None, compact=False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    header = table.rows[0]
    header._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    header._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    for index, label in enumerate(headers):
        cell = header.cells[index]
        set_cell_fill(cell, NAVY)
        set_cell_margins(cell, top=55 if compact else 100, bottom=55 if compact else 100)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        run = cell.paragraphs[0].add_run(label)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        if compact:
            run.font.size = Pt(9.5)
        if widths:
            cell.width = Inches(widths[index])
    for row_index, values in enumerate(rows):
        row = table.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        cells = row.cells
        for index, value in enumerate(values):
            cell = cells[index]
            set_cell_margins(cell, top=55 if compact else 100, bottom=55 if compact else 100)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_index % 2:
                set_cell_fill(cell, PALE)
            run = cell.paragraphs[0].add_run(str(value))
            if compact:
                run.font.size = Pt(9.5)
            if widths:
                cell.width = Inches(widths[index])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_code(doc, lines):
    paragraph = doc.add_paragraph()
    paragraph.style = doc.styles["No Spacing"]
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.right_indent = Inches(0.25)
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(9)
    properties = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F1F3F2")
    properties.append(shading)
    run = paragraph.add_run("\n".join(lines))
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    run.font.size = Pt(8.5)


def paragraph(doc, text, bold_lead=None):
    item = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        item.add_run(bold_lead).bold = True
        item.add_run(text[len(bold_lead):])
    else:
        item.add_run(text)
    return item


def bullet(doc, text, level=0):
    item = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    item.add_run(text)
    return item


def numbered(doc, text):
    item = doc.add_paragraph(style="List Number")
    item.add_run(text)
    return item


def section(doc, title, level=1):
    return doc.add_heading(title, level=level)


def build() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12

    for style_name, size in (("Title", 27), ("Heading 1", 18), ("Heading 2", 13), ("Heading 3", 11)):
        style = doc.styles[style_name]
        style.font.name = "Aptos Display" if style_name != "Heading 3" else "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(16 if style_name == "Heading 1" else 10)
        style.paragraph_format.space_after = Pt(7)
        if style_name == "Title":
            style_properties = style._element.get_or_add_pPr()
            style_borders = style_properties.find(qn("w:pBdr"))
            if style_borders is not None:
                style_properties.remove(style_borders)

    for section_item in doc.sections:
        section_item.top_margin = Inches(0.75)
        section_item.bottom_margin = Inches(0.7)
        section_item.left_margin = Inches(0.85)
        section_item.right_margin = Inches(0.85)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(92)
    title.add_run("Contract and Compliance Version Intelligence Platform")
    title_properties = title._p.get_or_add_pPr()
    title_borders = title_properties.find(qn("w:pBdr"))
    if title_borders is not None:
        title_properties.remove(title_borders)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_before = Pt(18)
    subtitle.add_run("Final Project Report").bold = True
    course = doc.add_paragraph()
    course.alignment = WD_ALIGN_PARAGRAPH.CENTER
    course.paragraph_format.space_before = Pt(38)
    course.add_run("Data Structures and Algorithms 3\n25CS2103E\nAcademic Year 2026 to 2027")
    team = doc.add_paragraph()
    team.alignment = WD_ALIGN_PARAGRAPH.CENTER
    team.paragraph_format.space_before = Pt(48)
    run = team.add_run("Team VAS\nTeam 05\nSection 13")
    run.bold = True
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_before = Pt(66)
    note.add_run("Local academic system using deterministic synthetic data").italic = True
    doc.add_page_break()

    first_section = doc.sections[0]
    first_section.different_first_page_header_footer = True
    header = first_section.header.paragraphs[0]
    header.text = "Contract and Compliance Version Intelligence Platform"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor(90, 101, 96)
    footer = first_section.footer.paragraphs[0]
    footer.add_run("Team VAS  |  ")
    add_page_number(footer)

    section(doc, "Executive Summary")
    paragraph(doc, "This project implements a local contract review and compliance workbench for a Data Structures and Algorithms course. The system connects a React and TypeScript interface to a FastAPI and SQLite backend. Six handwritten deterministic algorithms provide clause retrieval, version alignment, similarity analysis, graph clustering, compliance coverage, and reviewer assignment. The application stores contract lifecycle changes, review decisions, assignments, settings, and audit events without using external services or artificial intelligence.")
    paragraph(doc, "The delivered system is intended for demonstration and algorithm study. It uses synthetic contracts and produces explainable decision support. Its results are not legal advice, measured accuracy, or compliance certification.")

    section(doc, "Problem Definition")
    paragraph(doc, "Contract review requires users to find relevant clauses, compare versions, identify related agreements, determine whether obligations are covered, and allocate work to suitable reviewers. These tasks become difficult when records are spread across documents and changes cannot be traced. The project represents these tasks as explicit data structures and deterministic algorithms whose inputs, decisions, and outputs can be inspected.")

    section(doc, "Objectives and Scope")
    for text in [
        "Maintain contracts, versions, clauses, obligations, reviewers, assignments, review decisions, settings, and append-only audit events.",
        "Provide Boolean clause search backed by an inverted index.",
        "Compare clause sequences and changed wording with dynamic programming.",
        "Compute document similarity with TF-IDF and cosine similarity, then form clusters with graph traversal.",
        "Select clauses that cover obligations with a deterministic greedy Set Cover implementation.",
        "Propose reviewer assignments with a handwritten min-cost max-flow implementation.",
        "Present every major capability through a responsive frontend and a typed local API.",
    ]:
        bullet(doc, text)
    paragraph(doc, "Authentication, OCR, external data feeds, cloud storage, AI models, and production deployment are outside the project scope. Their absence is deliberate because they do not support the course objective of demonstrating the selected data structures and algorithms.")

    section(doc, "System Architecture")
    add_code(doc, [
        "React and TypeScript frontend",
        "        | typed HTTP requests",
        "        v",
        "FastAPI routes and domain services",
        "        |                    |",
        "        |                    +-- handwritten algorithms",
        "        v",
        "SQLAlchemy models and SQLite database",
    ])
    add_table(doc, ["Layer", "Technology", "Responsibility"], [
        ("Presentation", "React 19 and TypeScript", "Routes, forms, tables, charts, loading states, validation feedback, and responsive navigation"),
        ("Transport", "Typed fetch client", "One configurable API base URL, structured errors, request cancellation, and JSON or multipart requests"),
        ("Application", "FastAPI and Pydantic", "Typed endpoints, validation, pagination, mutation conflicts, and API documentation"),
        ("Domain", "Python modules", "Lifecycle rules, ingestion, identifiers, and six handwritten algorithms"),
        ("Persistence", "SQLAlchemy and SQLite", "Normalized records, deterministic seed data, and append-only audit events"),
    ], [1.15, 1.55, 3.7])

    data_model_heading = section(doc, "Data Model")
    data_model_heading.paragraph_format.page_break_before = True
    paragraph(doc, "Contracts own ordered versions. Versions contain clause records that retain stable clause keys across revisions. Tags and obligations use mapping tables so a clause can participate in several classifications and compliance controls. Reviewers have expertise values, workload, and capacity. Assignment, review, decision, settings, and audit records preserve workflow state.")
    add_table(doc, ["Entity", "Purpose", "Important relationships"], [
        ("Contract", "Agreement identity and lifecycle", "Has versions, review items, and assignment records"),
        ("Contract Version", "Immutable revision metadata", "Belongs to one contract and contains ordered clauses"),
        ("Clause", "Searchable provision with stable key", "Belongs to a version and maps to tags and obligations"),
        ("Obligation", "Compliance control", "Maps to zero or more clauses"),
        ("Reviewer", "Reviewer capability and capacity", "Has expertise values and assignments"),
        ("Review Decision", "Recorded workflow evidence", "Belongs to a review and stores notes and deciding user"),
        ("Audit Event", "Append-only mutation history", "References the affected entity and action"),
    ], [1.35, 2.15, 2.9])

    algorithms_heading = section(doc, "Handwritten Algorithms")
    algorithms_heading.paragraph_format.page_break_before = True
    add_table(doc, ["Capability", "Algorithm", "Main complexity"], [
        ("Clause retrieval", "Unicode-normalized inverted index with AND and OR postings", "Build O(T); query proportional to combined postings"),
        ("Version comparison", "Clause sequence LCS and word-level LCS", "O(nm) time and space"),
        ("Similarity", "TF-IDF and cosine similarity", "Pairwise worst case O(N^2 V)"),
        ("Clustering", "Threshold graph and iterative DFS", "O(V plus E) after graph construction"),
        ("Compliance", "Greedy Set Cover", "Straightforward worst case O(C O^2)"),
        ("Assignment", "Min-cost max-flow", "Successive shortest augmenting paths"),
    ], [1.35, 3.2, 1.85])

    search_heading = section(doc, "Inverted Index Search", 2)
    search_heading.paragraph_format.page_break_before = True
    paragraph(doc, "The index normalizes text with Unicode NFKC, applies case folding, tokenizes words, and records a sorted posting set for every term. An AND query intersects postings; an OR query unions them. Clause identifiers provide deterministic result ordering.")
    add_code(doc, ["for each clause:", "  for each unique normalized token:", "    postings[token].add(clause_id)", "result = intersection(postings[t] for t in query)  # AND"])

    section(doc, "Sequence Alignment", 2)
    paragraph(doc, "A dynamic-programming longest common subsequence table aligns stable clause keys between two versions. The reconstruction labels clauses as added, removed, unchanged, or modified. Modified clauses receive a second word-token LCS so the response can show specific wording evidence.")

    section(doc, "TF IDF Similarity and Clustering", 2)
    paragraph(doc, "Term frequency is the token count divided by document token count. Inverse document frequency uses log((N + 1) / (df + 1)) + 1. Sparse-vector cosine similarity supplies weighted edges for an undirected threshold graph. Iterative depth-first search returns sorted connected components.")

    section(doc, "Greedy Set Cover", 2)
    paragraph(doc, "At each step, the algorithm selects the clause that covers the largest number of currently uncovered obligations. Equal gains are resolved by the lexicographically smaller clause identifier. Obligations that no candidate can cover are returned explicitly instead of being hidden.")

    section(doc, "Reviewer Assignment", 2)
    paragraph(doc, "The residual network connects a source to contracts, valid contract-reviewer pairs, reviewers to a sink, and reverse edges for later correction. Reviewer-to-sink capacity equals remaining workload capacity. The objective is exactly: maximize valid assignments first, then minimize workload/expertise cost.")

    backend_heading = section(doc, "Backend API")
    backend_heading.paragraph_format.page_break_before = True
    paragraph(doc, "FastAPI exposes typed endpoints for health, dashboard summaries, contract and clause lifecycle operations, search, comparisons, graphs, coverage, reviewer management, assignment proposals and confirmation, review decisions, audit events, and settings. Lists use server pagination. Errors follow a consistent error object with code, message, and optional details.")
    add_table(doc, ["Operation", "Representative endpoint"], [
        ("Contracts and versions", "GET or POST /contracts and POST /contracts/{id}/versions"),
        ("Clause search", "POST /clause-search"),
        ("Version alignment", "POST /version-comparisons"),
        ("Similarity graph", "POST /similarity/graph"),
        ("Compliance coverage", "POST /compliance/coverage"),
        ("Assignment proposal", "POST /reviewer-assignments/propose"),
        ("Review decision", "POST /reviews/{id}/decision"),
        ("Settings", "GET and PUT /settings"),
    ], [2.3, 4.1])

    frontend_heading = section(doc, "Frontend Workbench")
    frontend_heading.paragraph_format.page_break_before = True
    paragraph(doc, "The frontend contains eleven routes: overview, contract register, contract detail, clause search, version analysis, similarity and clustering, compliance coverage, reviewer assignment, review queue, audit trail, and settings. API-backed views handle loading, error, empty, retry, and mutation states. Contract summaries are generated in the browser from already-loaded data.")
    paragraph(doc, "Responsive layouts collapse multi-column analysis views, expose a keyboard-accessible mobile navigation drawer, preserve table scrolling, and keep primary actions available on narrow screens. Dialogs support keyboard dismissal and restore focus to their trigger controls.")

    section(doc, "Deterministic Demonstration Data")
    add_table(doc, ["Record type", "Seeded count"], [
        ("Contracts", "12"), ("Versions", "36"), ("Clause-version rows", "276"),
        ("Obligations", "10"), ("Reviewers", "6"),
    ], [3.6, 2.8])
    paragraph(doc, "The seed command recreates the same identifiers, dates, relationships, and algorithm inputs. The corpus includes privacy, technology, procurement, commercial, research, and financial contract families together with added, removed, unchanged, and modified clauses.")

    section(doc, "Testing and Verification")
    add_table(doc, ["Test layer", "Coverage included"], [
        ("Backend pytest", "51 tests covering algorithms, edge cases, deterministic ties, persistence, lifecycle conflicts, validation, migrations, ingestion, and audit events"),
        ("Frontend Vitest", "6 component tests covering API payloads, selection persistence, lifecycle actions, obligation restoration, and review evidence"),
        ("Playwright", "6 browser workflows covering every route, record navigation, dialogs, mobile keyboard navigation, review evidence, and complete contract lifecycle operations"),
        ("Static checks", "Oxlint, TypeScript compilation, and the Vite production build"),
        ("Continuous integration", "Python 3.14.4 compatibility dry run, pinned dependency installation, backend tests, frontend checks, and Chromium browser tests"),
    ], [1.45, 4.95])

    section(doc, "Python Compatibility and Reproducibility")
    paragraph(doc, "The backend targets the existing Python 3.14.4 interpreter. Requirements pin FastAPI 0.141.1, SQLAlchemy 2.0.54, Pydantic 2.13.5, Uvicorn 0.53.0, pytest 9.1.1, and httpx 0.28.1. The documented setup performs a binary-only dependency dry run before installation. It does not install, replace, or downgrade Python.")

    section(doc, "Setup and Execution")
    add_code(doc, [
        "C:\\Python314\\python.exe -m venv .venv",
        ".\\.venv\\Scripts\\python.exe -m pip install -r backend\\requirements.txt",
        "npm install",
        "Set-Location backend",
        "..\\.venv\\Scripts\\python.exe -m app.seed reset",
        "..\\.venv\\Scripts\\python.exe -m uvicorn app.main:app --host 127.0.0.1 --port 8100",
        "# In a second terminal at the repository root",
        "npm run dev -- --host 127.0.0.1 --port 5181",
    ])

    section(doc, "Demonstration Procedure")
    for text in [
        "Reset the database and open the overview to establish the deterministic portfolio state.",
        "Search for a clause with AND and OR modes and explain the posting-list result evidence.",
        "Compare two versions and identify added, removed, and modified clauses.",
        "Change the similarity threshold and explain how graph edges alter connected components.",
        "Run compliance coverage and explain greedy gain selection and deterministic tie-breaking.",
        "Generate reviewer proposals, explain capacity constraints, and confirm assignments.",
        "Record a review decision with notes, then verify the corresponding audit event.",
        "Create a contract, add or import a version, archive it, and restore it.",
    ]:
        numbered(doc, text)

    section(doc, "Limitations")
    for text in [
        "The corpus is compact and synthetic, so the system does not claim measured legal accuracy.",
        "Text-bearing PDF, DOCX, and TXT ingestion is supported up to 10 MiB; scanned-only documents require OCR and are rejected.",
        "SQLite is suitable for the local demonstration but not for a concurrent production deployment.",
        "There is no authentication or role authorization because the application is local and academic.",
        "The algorithms favor explainability and deterministic behavior over large-corpus optimization.",
    ]:
        bullet(doc, text)

    section(doc, "Conclusion")
    paragraph(doc, "The project meets its academic objective by connecting six classical algorithms to persistent contract-review workflows. Users can inspect the source records, algorithm inputs, deterministic outputs, and audit history through one local application. The implementation remains intentionally bounded: it demonstrates data structures and algorithms without claiming legal authority or hiding results behind external AI services.")

    section(doc, "Viva Reference")
    add_table(doc, ["Question", "Answer focus"], [
        ("Why use an inverted index instead of scanning every clause?", "Posting lists avoid repeated full-corpus scans and make Boolean operations explicit."),
        ("Why LCS for version comparison?", "It preserves clause order and identifies a longest stable subsequence before classifying changes."),
        ("Why smooth IDF?", "The specified formula avoids division by zero and retains a positive weight for common terms."),
        ("Why DFS after thresholding?", "Connected components represent groups linked by similarity paths, not only direct pairs."),
        ("Is greedy Set Cover always optimal?", "No. It is deterministic and explainable but may not find the minimum global cover."),
        ("What does min-cost max-flow optimize?", "It maximizes the number of valid assignments before minimizing workload and expertise cost."),
        ("Why stable clause keys?", "They allow clause identity to survive across versions and support meaningful sequence alignment."),
        ("Why synthetic data?", "It produces repeatable demonstrations without exposing confidential agreements or fabricating real performance claims."),
    ], [2.9, 3.5], compact=True)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
